# OpenPyXL：读写 .xlsx（保留格式、公式、多工作表）
import os, tempfile
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.chart import BarChart, Reference
import openpyxl
print("openpyxl", openpyxl.__version__)

path = os.path.join(tempfile.gettempdir(), "report.xlsx")
wb = Workbook()
ws = wb.active
ws.title = "模型评测"
ws.append(["模型", "得分", "成本"])
for row in [("A", 85, 0.12), ("B", 72, 0.05), ("C", 91, 0.30)]:
    ws.append(row)
ws["E1"], ws["E2"] = "平均分", "=AVERAGE(B2:B4)"       # 写入的是公式字符串
for cell in ws[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="4472C4")
ws.column_dimensions["A"].width = 12

chart = BarChart(); chart.title = "得分"
chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=4), titles_from_data=True)
chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=4))
ws.add_chart(chart, "G2")
wb.save(path)

wb2 = load_workbook(path)
ws2 = wb2["模型评测"]
for row in ws2.iter_rows(min_row=2, max_row=4, values_only=True):
    print(row[:3])
print("E2 单元格内容:", ws2["E2"].value)
wb3 = load_workbook(path, data_only=True)
print("data_only=True 读到的值:", wb3["模型评测"]["E2"].value, "← openpyxl 不计算公式，需 Excel 等打开保存后才有缓存值")
